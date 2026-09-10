"""
Watchlist workbook: S3 access + shared sheet read/write helpers.

Used by both excel_source.py (the API's collection flow) and the standalone
stock_monitor.py script, so there's one place that knows the S3 bucket/key
and how to map/write the sheet's columns, instead of duplicating this logic
in both files. The workbook is never written to local disk — just held in
memory (BytesIO) for as long as it's being read or built.
"""

from datetime import datetime
from io import BytesIO

import openpyxl
from openpyxl.styles import Font

from app.core.aws_credentials import AwsCredentials
from app.core.config import settings

aws_credentials = AwsCredentials()

WATCHLIST_KEY = "My-watchlist-stocks.xlsx"

# Sheets that list the symbols we track.
WATCHLIST_SHEETS = ("Breakout Stocks CMP", "Buying Range Stocks CMP")

GREEN_BOLD = Font(color="008000", bold=True)
RED_BOLD = Font(color="FF0000", bold=True)
HEADER_ROW = 2  # row 1 is the title, row 2 has the real column names


# --- S3 access ---------------------------------------------------------


def download_watchlist() -> BytesIO:
    """Fetch the watchlist workbook from S3 into an in-memory buffer."""
    s3 = aws_credentials.s3Client()
    response = s3.get_object(Bucket=settings.S3_BUCKET_NAME, Key=WATCHLIST_KEY)
    return BytesIO(response["Body"].read())


def upload_watchlist(buffer: BytesIO) -> None:
    """Upload the (updated) watchlist workbook back to the same S3 key."""
    buffer.seek(0)
    s3 = aws_credentials.s3Client()
    s3.put_object(Bucket=settings.S3_BUCKET_NAME, Key=WATCHLIST_KEY, Body=buffer.read())


# --- Sheet read/write helpers -------------------------------------------


def build_column_map(worksheet) -> dict[str, int]:
    """Map header name -> column number, e.g. {"Symbol": 2, "CMP (₹)": 5, ...}"""
    column_of = {}
    for cell in worksheet[HEADER_ROW]:
        if cell.value:
            column_of[str(cell.value).strip()] = cell.column
    return column_of


def build_row_map(worksheet, symbol_column: int) -> dict[str, int]:
    """Map each symbol -> its row number, by scanning the Symbol column."""
    row_of_symbol = {}
    for row in range(HEADER_ROW + 1, worksheet.max_row + 1):
        cell_value = worksheet.cell(row=row, column=symbol_column).value
        if cell_value:
            row_of_symbol[str(cell_value).strip().upper()] = row
    return row_of_symbol


def read_watchlist_symbols() -> list[str]:
    """Return the de-duplicated, upper-cased symbol list from every watchlist sheet."""
    workbook = openpyxl.load_workbook(download_watchlist(), data_only=True)

    symbols: list[str] = []
    seen: set[str] = set()
    for sheet_name in WATCHLIST_SHEETS:
        if sheet_name not in workbook.sheetnames:
            continue
        worksheet = workbook[sheet_name]
        symbol_column = build_column_map(worksheet).get("Symbol")
        if symbol_column is None:
            continue
        for symbol in build_row_map(worksheet, symbol_column):
            if symbol not in seen:
                seen.add(symbol)
                symbols.append(symbol)
    return symbols


def write_common_fields(
    worksheet, row: int, column_of: dict[str, int], price, volume, sector, company_name
) -> None:
    """Write the columns every sheet shares: name, sector, CMP, volume, fetch status, updated at."""
    worksheet.cell(row=row, column=column_of["Company Name"], value=company_name)
    worksheet.cell(row=row, column=column_of["Sector"], value=sector)
    worksheet.cell(row=row, column=column_of["CMP (₹)"], value=price)
    worksheet.cell(row=row, column=column_of["Volume"], value=volume)
    worksheet.cell(row=row, column=column_of["Fetch Data"], value="Success")
    worksheet.cell(row=row, column=column_of["Update Time"], value=datetime.now())


def write_breakout_status(worksheet, row: int, column_of: dict[str, int], price, target) -> bool:
    """Write BreakOut Done/NA for this row. Returns whether it crossed."""
    crossed = target is not None and price is not None and price >= target

    cell = worksheet.cell(row=row, column=column_of["BreakOut"])
    if crossed:
        cell.value = "Done"
        cell.font = GREEN_BOLD
    else:
        cell.value = "NA"
        cell.font = Font()
    return crossed


def write_watching_level_status(
    worksheet, row: int, column_of: dict[str, int], price, watching_target
) -> bool:
    """Write Watching Level Reached/Not Reached for this row. Returns whether reached."""
    reached = False
    if watching_target is not None and price is not None and watching_target != 0:
        percent_gap = abs(price - watching_target) / watching_target * 100
        reached = percent_gap <= 2

    cell = worksheet.cell(row=row, column=column_of["Watching Level(1-2%)"])
    if reached:
        cell.value = "Reached"
        cell.font = GREEN_BOLD
    else:
        cell.value = "Not Reached"
        cell.font = RED_BOLD
    return reached
